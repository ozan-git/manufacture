# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).

import base64
import io
import re
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.modules.module import get_resource_path

try:  # pragma: no cover - optional dependency handled at runtime
    import openpyxl
except ImportError:  # pragma: no cover - guarded import
    openpyxl = None


class QcExportExcelWizard(models.TransientModel):
    _name = "qc.export.excel.wizard"
    _description = "Export quality tests to Excel"

    test_ids = fields.Many2many(
        comodel_name="qc.test",
        string="Tests",
        default=lambda self: self._default_test_ids(),
    )
    data_file = fields.Binary(string="File", readonly=True)
    filename = fields.Char(string="Filename")

    @api.model
    def _default_test_ids(self):
        active_ids = self.env.context.get("active_ids")
        if active_ids:
            return [(6, 0, active_ids)]
        return False

    def action_export(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(
                _(
                    "The python package 'openpyxl' is required to export Excel "
                    "files. Please install it on the server environment."
                )
            )
        tests = self.test_ids
        if not tests:
            raise UserError(_("Select at least one test to export."))
        workbook = self._load_template_workbook()
        try:
            sheet = workbook.active
            headers = self._read_template_headers(sheet)
            self._clear_template_rows(sheet)
        except Exception:
            workbook.close()
            raise
        for row in self._iter_template_rows(tests):
            sheet.append([row.get(column, "") for column in headers])
        self._populate_summary_sheet(workbook, tests)
        buffer = io.BytesIO()
        try:
            workbook.save(buffer)
        finally:
            workbook.close()
        buffer.seek(0)
        filename = self._build_filename(tests)
        self.write(
            {
                "data_file": base64.b64encode(buffer.read()),
                "filename": filename,
            }
        )
        return {
            "type": "ir.actions.act_url",
            "url": (
                "/web/content/?model=%s&id=%s&field=data_file&filename_field=filename"
                "&download=true"
            )
            % (self._name, self.id),
            "target": "self",
        }

    def _build_filename(self, tests):
        if len(tests) == 1:
            base = tests.code or tests.name or "qc_tests"
        else:
            base = "qc_tests"
        base = self._sanitize_filename_component(base)
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        return f"{base}_{timestamp}.xlsx"

    def _sanitize_filename_component(self, text):
        if not text:
            return "qc_tests"
        slug = re.sub(r"[^0-9A-Za-z-_]+", "_", text).strip("_")
        return slug or "qc_tests"

    def _iter_template_rows(self, tests):
        for test in tests.sorted(key=lambda t: (t.code or "", t.name or "")):
            product_payloads = self._prepare_product_payloads(test)
            if not product_payloads:
                product_payloads = [self._empty_product_payload()]
            for question in test.test_lines.sorted(
                key=lambda q: (q.sequence, q.id)
            ):
                base = self._prepare_base_row(test, question)
                for product_data in product_payloads:
                    if question.type == "qualitative":
                        values = question.ql_values or self.env[
                            "qc.test.question.value"
                        ]
                        if values:
                            for value in values.sorted(key=lambda v: v.id):
                                row = {**base, **product_data}
                                row.update(self._prepare_qualitative_payload(value))
                                yield row
                        else:
                            row = {**base, **product_data}
                            row.update(self._prepare_qualitative_payload(None))
                            yield row
                    else:
                        row = {**base, **product_data}
                        row.update(self._prepare_quantitative_payload(question))
                        yield row

    def _prepare_base_row(self, test, question):
        return {
            "name": test.name or "",
            "test_lines/name": question.name or "",
            "test_lines/type": question.type or "qualitative",
            "test_lines/notes": question.notes or "",
        }

    def _prepare_quantitative_payload(self, question):
        return {
            "test_lines/min_value": (
                question.min_value if question.min_value is not None else ""
            ),
            "test_lines/max_value": (
                question.max_value if question.max_value is not None else ""
            ),
            "test_lines/uom": question.uom_id.display_name
            if question.uom_id
            else "",
            "test_lines/ql_values/name": "",
            "test_lines/ql_values/ok": "",
        }

    def _prepare_qualitative_payload(self, value):
        if not value:
            return {
                "test_lines/min_value": "",
                "test_lines/max_value": "",
                "test_lines/uom": "",
                "test_lines/ql_values/name": "",
                "test_lines/ql_values/ok": False,
            }
        return {
            "test_lines/min_value": "",
            "test_lines/max_value": "",
            "test_lines/uom": "",
            "test_lines/ql_values/name": value.name or "",
            "test_lines/ql_values/ok": bool(value.ok),
        }

    def _prepare_product_payloads(self, test):
        payloads = []
        seen = set()

        def add_payload(template, trigger_name, trigger_timing):
            payload = self._build_product_payload(template, trigger_name, trigger_timing)
            key = (
                payload[
                    "trigger_product_template_line_ids/product_template/default_code"
                ],
                payload["trigger_product_template_line_ids/product_template/name"],
                payload["trigger_product_template_line_ids/trigger/name"],
                payload["trigger_product_template_line_ids/timing"],
            )
            if key in seen:
                return
            seen.add(key)
            payloads.append(payload)

        TemplateLine = self.env["qc.trigger.product_template_line"]
        for trigger_line in TemplateLine.search([("test", "=", test.id)]):
            add_payload(
                trigger_line.product_template,
                trigger_line.trigger.name,
                trigger_line.timing,
            )

        ProductLine = self.env["qc.trigger.product_line"]
        for trigger_line in ProductLine.search([("test", "=", test.id)]):
            template = (
                trigger_line.product.product_tmpl_id
                if trigger_line.product
                else False
            )
            add_payload(template, trigger_line.trigger.name, trigger_line.timing)

        CategoryLine = self.env["qc.trigger.product_category_line"]
        for trigger_line in CategoryLine.search([("test", "=", test.id)]):
            add_payload(False, trigger_line.trigger.name, trigger_line.timing)

        related_template = self._resolve_related_product_template(test)
        if related_template:
            add_payload(related_template, "", "")

        payloads.sort(
            key=lambda payload: (
                payload[
                    "trigger_product_template_line_ids/product_template/default_code"
                ],
                payload["trigger_product_template_line_ids/product_template/name"],
                payload["trigger_product_template_line_ids/trigger/name"],
                payload["trigger_product_template_line_ids/timing"],
            )
        )
        return payloads

    def _build_product_payload(self, template, trigger_name, trigger_timing):
        default_code = template.default_code if template else ""
        name = template.display_name if template else ""
        return {
            "trigger_product_template_line_ids/product_template/default_code": (
                default_code or ""
            ),
            "trigger_product_template_line_ids/product_template/name": name or "",
            "trigger_product_template_line_ids/trigger/name": trigger_name or "",
            "trigger_product_template_line_ids/timing": (
                trigger_timing or ("after" if trigger_name else "")
            ),
        }

    def _empty_product_payload(self):
        return {
            "trigger_product_template_line_ids/product_template/default_code": "",
            "trigger_product_template_line_ids/product_template/name": "",
            "trigger_product_template_line_ids/trigger/name": "",
            "trigger_product_template_line_ids/timing": "",
        }

    def _resolve_related_product_template(self, test):
        if getattr(test, "type", False) != "related" or not test.object_id:
            return False
        reference = test.object_id
        if hasattr(reference, "product_tmpl_id") and reference.product_tmpl_id:
            return reference.product_tmpl_id
        if hasattr(reference, "product_id") and reference.product_id:
            return reference.product_id.product_tmpl_id
        if getattr(reference, "_name", "") == "product.template":
            return reference
        return False

    def _get_external_id(self, record):
        if not record:
            return ""
        xmlids = record.get_external_id()
        return xmlids.get(record.id, "")

    def _clear_template_rows(self, sheet):
        if sheet.max_row and sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)

    def _get_template_headers(self):
        workbook = self._load_template_workbook(read_only=True)
        try:
            return self._read_template_headers(workbook.active)
        finally:
            workbook.close()

    def _load_template_workbook(self, read_only=False):
        if openpyxl is None:
            raise UserError(
                _(
                    "The python package 'openpyxl' is required to load the Excel "
                    "template. Please install it on the server environment."
                )
            )
        template_path = self._get_template_path()
        try:
            return openpyxl.load_workbook(template_path, read_only=read_only)
        except Exception as exc:
            raise UserError(
                _(
                    "The quality control Excel template could not be loaded. "
                    "Please reinstall the module or restore the original template."
                )
            ) from exc

    def _read_template_headers(self, sheet):
        try:
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        except StopIteration as exc:
            raise UserError(
                _(
                    "The quality control Excel template is missing the header "
                    "row. Please reinstall the module or restore the original "
                    "template."
                )
            ) from exc
        return [value or "" for value in first_row]

    def _get_template_path(self):
        template_path = get_resource_path(
            "quality_control_oca",
            "static/xlsx",
            "qc_product_questions_template.xlsx",
        )
        if not template_path:
            raise UserError(
                _(
                    "The quality control Excel template could not be located in "
                    "the module resources."
                )
            )
        return template_path

    def _populate_summary_sheet(self, workbook, tests):
        """Add a human-readable sheet with the exported details."""
        summary_columns = [
            (_("Test Code"), "test_code"),
            (_("Test Name"), "test_name"),
            (_("Test Category"), "test_category"),
            (_("Test Type"), "test_type"),
            (_("Reference Object"), "reference_object"),
            (_("Fill Correct Values"), "fill_correct_values"),
            (_("Product Default Code"), "product_default_code"),
            (_("Product Name"), "product_name"),
            (_("Trigger"), "trigger_name"),
            (_("Trigger Timing"), "trigger_timing"),
            (_("Question Sequence"), "question_sequence"),
            (_("Question Code"), "question_code"),
            (_("Question Name"), "question_name"),
            (_("Question Type"), "question_type"),
            (_("Question Notes"), "question_notes"),
            (_("Question Unit of Measure"), "question_uom"),
            (_("Question Min Value"), "question_min_value"),
            (_("Question Max Value"), "question_max_value"),
            (_("Option"), "option_name"),
            (_("Option Is Correct"), "option_is_correct"),
            (_("Row Summary"), "summary"),
        ]
        summary_sheet = workbook.create_sheet(title=_("Questions Summary"))
        summary_sheet.append([label for label, _ in summary_columns])
        for row in self._iter_summary_rows(tests):
            summary_sheet.append(
                [row.get(column_key, "") for _, column_key in summary_columns]
            )

    def _iter_summary_rows(self, tests):
        for test in tests.sorted(key=lambda t: (t.code or "", t.name or "")):
            product_payloads = self._prepare_product_payloads(test)
            if not product_payloads:
                product_payloads = [self._empty_product_payload()]
            base_values = self._prepare_summary_base_values(test)
            for question in test.test_lines.sorted(key=lambda q: (q.sequence, q.id)):
                question_values = self._prepare_summary_question(question)
                for product_data in product_payloads:
                    product_values = self._prepare_summary_product(product_data)
                    combined = {**base_values, **question_values, **product_values}
                    if question.type == "qualitative":
                        values = question.ql_values.sorted(key=lambda v: v.id)
                        if values:
                            for value in values:
                                payload = {
                                    **combined,
                                    **self._prepare_summary_qualitative_value(value),
                                }
                                payload["summary"] = self._build_summary_label(payload)
                                yield payload
                        else:
                            payload = {
                                **combined,
                                **self._prepare_summary_qualitative_value(None),
                            }
                            payload["summary"] = self._build_summary_label(payload)
                            yield payload
                    else:
                        payload = {
                            **combined,
                            **self._prepare_summary_quantitative_values(question),
                        }
                        payload["summary"] = self._build_summary_label(payload)
                        yield payload

    def _prepare_summary_base_values(self, test):
        return {
            "test_code": test.code or "",
            "test_name": test.name or "",
            "test_category": test.category.display_name if test.category else "",
            "test_type": test.type or "generic",
            "reference_object": test.object_id.display_name
            if test.object_id
            else "",
            "fill_correct_values": self._format_boolean(test.fill_correct_values),
        }

    def _prepare_summary_product(self, product_payload):
        return {
            "product_default_code": product_payload.get(
                "trigger_product_template_line_ids/product_template/default_code", ""
            ),
            "product_name": product_payload.get(
                "trigger_product_template_line_ids/product_template/name", ""
            ),
            "trigger_name": product_payload.get(
                "trigger_product_template_line_ids/trigger/name", ""
            ),
            "trigger_timing": product_payload.get(
                "trigger_product_template_line_ids/timing", ""
            ),
        }

    def _prepare_summary_question(self, question):
        return {
            "question_sequence": question.sequence or 0,
            "question_code": question.code or "",
            "question_name": question.name or "",
            "question_type": question.type or "qualitative",
            "question_notes": question.notes or "",
            "question_uom": question.uom_id.display_name if question.uom_id else "",
            "question_min_value": question.min_value
            if question.min_value is not None
            else "",
            "question_max_value": question.max_value
            if question.max_value is not None
            else "",
        }

    def _prepare_summary_quantitative_values(self, question):
        return {
            "option_name": "",
            "option_is_correct": "",
            "question_min_value": question.min_value
            if question.min_value is not None
            else "",
            "question_max_value": question.max_value
            if question.max_value is not None
            else "",
        }

    def _prepare_summary_qualitative_value(self, value):
        if not value:
            return {
                "option_name": "",
                "option_is_correct": "",
                "question_min_value": "",
                "question_max_value": "",
            }
        return {
            "option_name": value.name or "",
            "option_is_correct": self._format_boolean(value.ok, empty_when_none=True),
            "question_min_value": "",
            "question_max_value": "",
        }

    def _format_boolean(self, value, empty_when_none=False):
        if value is None and empty_when_none:
            return ""
        return _("Yes") if value else _("No")

    def _build_summary_label(self, payload):
        parts = [
            payload.get("test_name") or "",
            payload.get("question_name") or "",
        ]
        option = payload.get("option_name")
        if option:
            parts.append(option)
        return " / ".join(part for part in parts if part)
