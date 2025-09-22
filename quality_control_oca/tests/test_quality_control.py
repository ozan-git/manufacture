# Copyright 2010 NaN Projectes de Programari Lliure, S.L.
# Copyright 2014 Serv. Tec. Avanzados - Pedro M. Baeza
# Copyright 2014 Oihane Crucelaegui - AvanzOSC
# Copyright 2017 ForgeFlow S.L.
# Copyright 2017 Simone Rubino - Agile Business Group
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).

import base64
import io

from odoo import exceptions
from odoo.tests import new_test_user

from odoo.addons.base.models.ir_model import MODULE_UNINSTALL_FLAG
from odoo.addons.base.tests.common import BaseCommon

from ..models.qc_trigger_line import _filter_trigger_lines

try:  # pragma: no cover - optional dependency handled at runtime
    import openpyxl
except ImportError:  # pragma: no cover - guarded import
    openpyxl = None


class TestQualityControlOcaBase(BaseCommon):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.inspection_model = cls.env["qc.inspection"]
        cls.category_model = cls.env["qc.test.category"]
        cls.question_model = cls.env["qc.test.question"]
        cls.wizard_model = cls.env["qc.inspection.set.test"]
        cls.qc_trigger = cls.env["qc.trigger"].create({"name": "Test Trigger"})
        cls.test = cls.env.ref("quality_control_oca.qc_test_1")
        cls.val_ok = cls.env.ref("quality_control_oca.qc_test_question_value_1")
        cls.val_ko = cls.env.ref("quality_control_oca.qc_test_question_value_2")
        cls.qn_question = cls.env.ref("quality_control_oca.qc_test_question_2")
        cls.cat_generic = cls.env.ref(
            "quality_control_oca.qc_test_template_category_generic"
        )
        cls.product = cls.env["product.product"].create({"name": "Test product"})
        cls.inspection1 = cls.inspection_model.create(
            {
                "name": "Test Inspection",
                "inspection_lines": cls.inspection_model._prepare_inspection_lines(
                    cls.test
                ),
            }
        )
        cls.user = new_test_user(
            cls.env,
            login="test_quality_control_oca",
            groups="quality_control_oca.group_quality_control_user",
        )


class TestQualityControlOca(TestQualityControlOcaBase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.wizard = cls.wizard_model.with_context(active_id=cls.inspection1.id).create(
            {"test": cls.test.id}
        )
        cls.wizard.action_create_test()
        cls.inspection1.action_todo()

    def test_inspection_correct(self):
        for line in self.inspection1.inspection_lines:
            if line.question_type == "qualitative":
                line.qualitative_value = self.val_ok
            if line.question_type == "quantitative":
                line.quantitative_value = 5.0
        self.inspection1.action_confirm()
        for line in self.inspection1.inspection_lines:
            self.assertTrue(
                line.success, f"Incorrect state in inspection line {line.name}"
            )

        self.assertTrue(
            self.inspection1.success,
            f"Incorrect state in inspection {self.inspection1.name}",
        )

        self.assertEqual(self.inspection1.state, "success")
        self.inspection1.action_approve()
        self.assertEqual(self.inspection1.state, "success")
        self.assertTrue(bool(self.inspection1.date_done))
        self.inspection1.action_cancel()
        self.inspection1.action_draft()
        self.assertFalse(self.inspection1.date_done)

    def test_inspection_incorrect(self):
        for line in self.inspection1.inspection_lines:
            if line.question_type == "qualitative":
                line.qualitative_value = self.val_ko
            if line.question_type == "quantitative":
                line.quantitative_value = 15.0
        self.inspection1.action_confirm()
        for line in self.inspection1.inspection_lines:
            self.assertFalse(
                line.success, f"Incorrect state in inspection line {line.name}"
            )
        self.assertFalse(
            self.inspection1.success,
            f"Incorrect state in inspection {self.inspection1.name}",
        )

        self.assertEqual(self.inspection1.state, "waiting")
        self.inspection1.action_approve()
        self.assertEqual(self.inspection1.state, "failed")
        self.assertTrue(bool(self.inspection1.date_done))

    def test_actions_errors(self):
        inspection2 = self.inspection1.copy()
        inspection2.action_draft()
        inspection2.write({"test": False})
        with self.assertRaises(exceptions.UserError):
            inspection2.action_todo()
        inspection3 = self.inspection1.copy()
        inspection3.write(
            {
                "inspection_lines": self.inspection_model._prepare_inspection_lines(
                    inspection3.test
                )
            }
        )
        for line in inspection3.inspection_lines:
            if line.question_type == "quantitative":
                line.quantitative_value = 15.0
        with self.assertRaises(exceptions.UserError):
            inspection3.action_confirm()
        inspection4 = self.inspection1.copy()
        inspection4.write(
            {
                "inspection_lines": self.inspection_model._prepare_inspection_lines(
                    inspection4.test
                )
            }
        )
        for line in inspection4.inspection_lines:
            if line.question_type == "quantitative":
                line.write({"uom_id": False, "quantitative_value": 15.0})
            elif line.question_type == "qualitative":
                line.qualitative_value = self.val_ok
        with self.assertRaises(exceptions.UserError):
            inspection4.action_confirm()

    def test_import_template_available(self):
        templates = self.env["qc.test"].get_import_templates()
        template_url = "/quality_control_oca/static/xlsx/qc_product_questions_template.xlsx"
        self.assertTrue(
            any(template.get("template") == template_url for template in templates),
            "The quality test import template should be exposed to the generic import view.",
        )

    def test_export_contains_template_headers(self):
        if openpyxl is None:
            self.skipTest("openpyxl not installed")
        wizard = (
            self.env["qc.export.excel.wizard"].with_context(active_ids=[self.test.id]).create({})
        )
        wizard.action_export()
        self.assertTrue(wizard.data_file, "The export wizard should generate a file.")
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.data_file)))
        try:
            sheet = workbook.active
            headers = list(
                next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            )
        finally:
            workbook.close()
        self.assertListEqual(headers, wizard._get_template_headers())

    def test_export_rows_follow_template_structure(self):
        if openpyxl is None:
            self.skipTest("openpyxl not installed")
        wizard = (
            self.env["qc.export.excel.wizard"].with_context(active_ids=[self.test.id]).create({})
        )
        wizard.action_export()
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.data_file)))
        try:
            sheet = workbook.active
            headers = list(
                next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            )
            rows = [
                list(row)
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if any(row)
            ]
        finally:
            workbook.close()

        self.assertTrue(rows, "The export should include at least one data row.")
        header_index = {name: index for index, name in enumerate(headers)}

        test_names = {
            row[header_index["name"]]
            for row in rows
            if row[header_index["name"]]
        }
        self.assertIn(self.test.name, test_names)

        question_types = {row[header_index["test_lines/type"]] for row in rows}
        self.assertIn("qualitative", question_types)
        self.assertIn("quantitative", question_types)

        qualitative_names = {
            row[header_index["test_lines/ql_values/name"]]
            for row in rows
            if row[header_index["test_lines/ql_values/name"]]
        }
        self.assertIn(self.val_ok.name, qualitative_names)
        self.assertIn(self.val_ko.name, qualitative_names)

        uom_xmlid = self.qn_question.uom_id.get_external_id().get(
            self.qn_question.uom_id.id
        )
        self.assertTrue(uom_xmlid, "The unit of measure should have a stable external ID.")
        quantitative_rows = [
            row
            for row in rows
            if row[header_index["test_lines/type"]] == "quantitative"
        ]
        uom_values = {
            row[header_index["test_lines/uom_id/id"]] for row in quantitative_rows
        }
        self.assertIn(uom_xmlid, uom_values)

    def test_export_includes_trigger_and_product_details(self):
        if openpyxl is None:
            self.skipTest("openpyxl not installed")

        product_template = self.env["product.template"].create(
            {"name": "Export Template", "default_code": "EXP-TPL"}
        )
        uom_unit = self.env.ref("uom.product_uom_unit")
        export_test = self.env["qc.test"].create(
            {
                "name": "Exported Test",
                "code": "EXP-TEST",
                "category": self.cat_generic.id,
                "fill_correct_values": True,
                "test_lines": [
                    (
                        0,
                        0,
                        {
                            "name": "Length control",
                            "code": "LEN",
                            "type": "quantitative",
                            "sequence": 5,
                            "min_value": 1.0,
                            "max_value": 5.0,
                            "uom_id": uom_unit.id,
                        },
                    )
                ],
            }
        )
        self.env["qc.trigger.product_template_line"].create(
            {
                "trigger": self.qc_trigger.id,
                "test": export_test.id,
                "product_template": product_template.id,
                "timing": "before",
            }
        )

        wizard = (
            self.env["qc.export.excel.wizard"]
            .with_context(active_ids=[export_test.id])
            .create({})
        )
        wizard.action_export()
        workbook = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(wizard.data_file))
        )
        try:
            sheet = workbook.active
            headers = list(
                next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            )
            header_index = {name: index for index, name in enumerate(headers)}
            rows = [
                list(row)
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if any(row)
            ]
        finally:
            workbook.close()

        self.assertTrue(rows, "The export should include at least one data row.")
        export_row = next(
            (
                row
                for row in rows
                if row[header_index["code"]] == export_test.code
            ),
            None,
        )
        self.assertIsNotNone(export_row, "The exported test should be present in the sheet.")

        category_xmlid = self.cat_generic.get_external_id().get(self.cat_generic.id)
        self.assertEqual(
            export_row[
                header_index[
                    "trigger_product_template_line_ids/product_template/default_code"
                ]
            ],
            product_template.default_code,
        )
        self.assertEqual(
            export_row[
                header_index[
                    "trigger_product_template_line_ids/product_template/name"
                ]
            ],
            product_template.display_name,
        )
        self.assertEqual(
            export_row[
                header_index["trigger_product_template_line_ids/trigger/name"]
            ],
            self.qc_trigger.name,
        )
        self.assertEqual(
            export_row[header_index["trigger_product_template_line_ids/timing"]],
            "before",
        )
        self.assertEqual(export_row[header_index["code"]], export_test.code)
        self.assertEqual(export_row[header_index["name"]], export_test.name)
        self.assertEqual(export_row[header_index["type"]], export_test.type)
        self.assertEqual(export_row[header_index["category/id"]], category_xmlid)
        self.assertTrue(export_row[header_index["fill_correct_values"]])

        question = export_test.test_lines
        self.assertEqual(
            export_row[header_index["test_lines/sequence"]], question.sequence
        )
        self.assertEqual(export_row[header_index["test_lines/code"]], question.code)
        self.assertEqual(export_row[header_index["test_lines/name"]], question.name)
        self.assertEqual(export_row[header_index["test_lines/type"]], question.type)
        self.assertIn(
            export_row[header_index["test_lines/notes"]],
            (question.notes, None, ""),
        )
        self.assertEqual(
            export_row[header_index["test_lines/uom_id/id"]],
            uom_unit.get_external_id().get(uom_unit.id),
        )
        self.assertEqual(
            export_row[header_index["test_lines/min_value"]], question.min_value
        )
        self.assertEqual(
            export_row[header_index["test_lines/max_value"]], question.max_value
        )
        self.assertFalse(export_row[header_index["test_lines/ql_values/name"]])
        self.assertFalse(export_row[header_index["test_lines/ql_values/ok"]])

    def test_categories(self):
        category1 = self.category_model.create({"name": "Category ONE"})
        category2 = self.category_model.create(
            {"name": "Category TWO", "parent_id": category1.id}
        )
        self.assertEqual(
            category2.complete_name,
            f"{category1.name} / {category2.name}",
            "Something went wrong when computing complete name",
        )
        with self.assertRaises(exceptions.UserError):
            category1.parent_id = category2.id

    def test_get_qc_trigger_product(self):
        self.test.write({"fill_correct_values": True})
        trigger_lines = set()
        self.product.write(
            {
                "qc_triggers": [
                    (0, 0, {"trigger": self.qc_trigger.id, "test": self.test.id})
                ],
            }
        )
        self.product.product_tmpl_id.write(
            {
                "qc_triggers": [
                    (0, 0, {"trigger": self.qc_trigger.id, "test": self.test.id})
                ],
            }
        )
        self.product.categ_id.write(
            {
                "qc_triggers": [
                    (0, 0, {"trigger": self.qc_trigger.id, "test": self.test.id})
                ],
            }
        )
        for model in [
            "qc.trigger.product_category_line",
            "qc.trigger.product_template_line",
            "qc.trigger.product_line",
        ]:
            trigger_lines = trigger_lines.union(
                self.env[model].get_trigger_line_for_product(
                    self.qc_trigger, ["after"], self.product
                )
            )
        self.assertEqual(len(trigger_lines), 3)
        filtered_trigger_lines = _filter_trigger_lines(trigger_lines)
        self.assertEqual(len(filtered_trigger_lines), 1)
        for trigger_line in filtered_trigger_lines:
            inspection = self.inspection_model._make_inspection(
                self.product, trigger_line
            )
            self.assertEqual(inspection.state, "ready")
            self.assertTrue(inspection.auto_generated)
            self.assertEqual(inspection.test, self.test)
            for line in inspection.inspection_lines:
                if line.question_type == "qualitative":
                    self.assertEqual(line.qualitative_value, self.val_ok)
                elif line.question_type == "quantitative":
                    self.assertAlmostEqual(
                        round(line.quantitative_value, 2),
                        round(
                            (self.qn_question.min_value + self.qn_question.max_value)
                            * 0.5,
                            2,
                        ),
                    )

    def test_qc_inspection_not_draft_unlink(self):
        with self.assertRaises(exceptions.UserError):
            self.inspection1.unlink()
        inspection2 = self.inspection1.copy()
        inspection2.action_cancel()
        self.assertEqual(inspection2.state, "canceled")
        inspection2.action_draft()
        self.assertEqual(inspection2.state, "draft")
        inspection2.unlink()

    def test_qc_inspection_auto_generate_manual_unlink(self):
        inspection2 = self.inspection1.copy()
        inspection2.write({"auto_generated": True})
        with self.assertRaises(exceptions.UserError):
            inspection2.with_user(self.user).unlink()
        self.assertTrue(inspection2.unlink())

    def test_qc_inspection_auto_generate_uninstall_unlink(self):
        uninstall = {MODULE_UNINSTALL_FLAG: True}

        inspection2 = self.inspection1.copy()
        inspection2.write({"auto_generated": True})
        self.assertTrue(inspection2.with_context(**uninstall).unlink())

    def test_qc_inspection_product(self):
        self.inspection1.write(
            {"object_id": "%s,%d" % (self.product._name, self.product.id)}
        )
        self.assertEqual(self.inspection1.product_id, self.product)

    def test_qc_test_question_constraints(self):
        with self.assertRaises(exceptions.ValidationError):
            self.question_model.create(
                {
                    "name": "Quantitative Question",
                    "type": "quantitative",
                    "min_value": 1.0,
                    "max_value": 0.0,
                }
            )
        with self.assertRaises(exceptions.ValidationError):
            self.question_model.create(
                {
                    "name": "Qualitative Question",
                    "type": "qualitative",
                    "ql_values": [(0, 0, {"name": "Qualitative answer", "ok": False})],
                }
            )

    def test_legacy_manager_can_delete_restricted_inspections(self):
        legacy_manager = new_test_user(
            self.env,
            login="test_quality_control_legacy_manager",
            groups="quality_control.group_quality_control_manager",
        )

        auto_generated_inspection = self.inspection_model.create(
            {
                "name": "Auto-generated inspection",
                "auto_generated": True,
            }
        )
        non_draft_inspection = self.inspection_model.create(
            {
                "name": "Non draft inspection",
                "state": "ready",
            }
        )

        auto_generated_inspection.with_user(legacy_manager).unlink()
        self.assertFalse(auto_generated_inspection.exists())

        non_draft_inspection.with_user(legacy_manager).unlink()
        self.assertFalse(non_draft_inspection.exists())
